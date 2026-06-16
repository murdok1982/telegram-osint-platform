import os
import datetime
import logging

logger = logging.getLogger(__name__)


def generate_docx_report(case_data: dict, output_path: str):
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        logger.error("python-docx not installed. Run: pip install python-docx")
        raise

    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)

    doc = Document()

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    title = doc.add_heading(f"CTI Intelligence Report - Case #{case_data['id']}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.add_paragraph()

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(case_data.get('summary', 'No summary available.'))

    doc.add_heading("2. Risk Assessment", level=1)
    table = doc.add_table(rows=4, cols=2)
    table.style = 'Light Grid Accent 1'

    cells = table.rows[0].cells
    cells[0].text = "Metric"
    cells[1].text = "Value"

    cells = table.rows[1].cells
    cells[0].text = "Risk Score"
    cells[1].text = f"{case_data.get('risk_score', 0)}/100"

    cells = table.rows[2].cells
    cells[0].text = "Status"
    cells[1].text = case_data.get('status', 'N/A')

    cells = table.rows[3].cells
    cells[0].text = "User Hash"
    cells[1].text = case_data.get('user_hash', 'N/A')

    doc.add_paragraph()

    doc.add_heading("3. Specialized Analyses", level=1)

    analysis = case_data.get('agent_analysis', {})

    doc.add_heading("Agent A - Psychological Profile", level=2)
    agent_a = analysis.get('agent_a', 'Pending')
    if isinstance(agent_a, str):
        doc.add_paragraph(agent_a)
    else:
        doc.add_paragraph(str(agent_a))

    doc.add_heading("Agent B - Threat Intelligence", level=2)
    agent_b = analysis.get('agent_b', 'Pending')
    if isinstance(agent_b, str):
        doc.add_paragraph(agent_b)
    else:
        doc.add_paragraph(str(agent_b))

    doc.add_heading("Agent C - Strategic Consolidation", level=2)
    agent_c = analysis.get('agent_c', 'Pending')
    if isinstance(agent_c, str):
        doc.add_paragraph(agent_c)
    else:
        doc.add_paragraph(str(agent_c))

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("--- CONFIDENTIAL --- Telegram OSINT/CTI Platform ---")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(128, 128, 128)

    doc.save(output_path)
    logger.info(f"DOCX report generated: {output_path}")
    return output_path
